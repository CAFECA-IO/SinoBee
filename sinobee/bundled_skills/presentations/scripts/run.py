#!/usr/bin/env python3
"""Presentations (.pptx): build from a JSON deck, read, check, render.

    run.py build <deck.json> <out.pptx>    16:9 deck drawn on blank slides with one theme: title, section,
                                           bullets, two-column, image, chart (native, editable), table, quote,
                                           big-number and closing slides; speaker notes; CJK fonts
    run.py inspect <file.pptx>             slides with title, text, tables, images, charts, notes
    run.py check <file.pptx>               text that overflows its box, shapes off the slide, empty
                                           placeholders, tiny fonts, crowded slides, slides without a title
    run.py render <file.pptx> <outdir>     PNG of every slide via LibreOffice + pdftoppm, to look at

Every command prints JSON. Exit 0 = ok, 1 = check found problems, 2 = usage / input error.
"""
from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import sys
from pathlib import Path

CJK = re.compile(r"[⺀-鿿가-힯豈-﫿＀-￯]")
W_IN, H_IN = 13.333, 7.5
EAST = {"zh-Hant": "Microsoft JhengHei", "zh-Hans": "Microsoft YaHei", "ja": "Yu Gothic", "ko": "Malgun Gothic"}
PLACEHOLDER = re.compile(r"lorem ipsum|click to add|按一下以新增|\bTODO\b|\bTBD\b|替換成|請填入", re.I)


def out(data: dict, code: int = 0):
    print(json.dumps(data, ensure_ascii=False, indent=1))
    sys.exit(code)


def need():
    try:
        import pptx  # noqa: F401
    except ImportError:
        out({"status": "error", "error": "python-pptx is not installed in this Python",
             "fix": f"{sys.executable} -m pip install python-pptx"}, 2)


def rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h.lstrip("#").upper())


# ------------------------------------------------------------------ build

class Deck:
    def __init__(self, spec, base: Path):
        from pptx import Presentation
        from pptx.util import Inches
        self.spec, self.base = spec, base
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W_IN), Inches(H_IN)
        th = spec.get("theme", {})
        text = json.dumps(spec, ensure_ascii=False)
        lang = spec.get("lang") or ("zh-Hant" if CJK.search(text) else "en")
        self.lang = lang
        self.font = th.get("font", "Calibri")
        self.east = th.get("east_asia") or EAST.get(lang, "Microsoft JhengHei")
        self.accent = th.get("accent", "1F4E79")
        self.accent2 = th.get("accent2", "C55A11")
        self.bg = th.get("background", "FFFFFF")
        self.fg = th.get("text", "222222")
        self.muted = th.get("muted", "6B6B6B")
        self.n = 0

    # -- primitives
    def slide(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(self.bg)
        self.n += 1
        return s

    def text(self, s, x, y, w, h, paras, size=18, bold=False, color=None, align="left", anchor="top",
             bullets=False, line=1.15):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for side in ("margin_left", "margin_right"):
            setattr(tf, side, Inches(0.05))
        items = paras if isinstance(paras, list) else [paras]
        for i, item in enumerate(items):
            level, t = 0, item
            if isinstance(item, dict):
                level, t = int(item.get("level", 0)), item.get("text", "")
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line
            if bullets:
                p.space_before = Pt(8 if level == 0 else 3)
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(0.3 + 0.35 * level))))
                pPr.set("indent", str(int(-Inches(0.25))))
                bu = pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"})
                pPr.append(bu)
            self._runs(p, str(t), size - (2 * level if bullets else 0), bold, color or self.fg)
        return box

    def _runs(self, p, t, size, bold, color):
        from pptx.util import Pt
        pos = 0
        for m in list(re.finditer(r"\*\*(.+?)\*\*", t)) + [None]:
            seg_end = m.start() if m else len(t)
            for seg, b in ((t[pos:seg_end], bold), ((m.group(1), True) if m else ("", bold))):
                if not seg:
                    continue
                r = p.add_run()
                r.text = seg
                f = r.font
                f.size, f.bold, f.name = Pt(size), b, self.font
                f.color.rgb = rgb(color)
                rpr = r._r.get_or_add_rPr()
                ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea",
                                     {"typeface": self.east})
                rpr.append(ea)
            if m:
                pos = m.end()

    def bar(self, s, x, y, w, h, color):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(color)
        r.line.fill.background()
        return r

    def title(self, s, t, sub=None):
        self.bar(s, 0.6, 0.55, 0.12, 0.62, self.accent)
        self.text(s, 0.85, 0.42, 11.8, 0.9, t, size=30, bold=True, anchor="middle")
        if sub:
            self.text(s, 0.85, 1.25, 11.8, 0.5, sub, size=16, color=self.muted)

    def footer(self, s):
        if self.spec.get("footer", True) is False:
            return
        label = self.spec.get("footer") if isinstance(self.spec.get("footer"), str) else ""
        self.text(s, 0.6, 7.0, 9, 0.35, label, size=10, color=self.muted)
        self.text(s, 11.7, 7.0, 1.0, 0.35, str(self.n), size=10, color=self.muted, align="right")

    def notes(self, s, sl):
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    def image(self, s, path, x, y, w, h):
        from pptx.util import Inches
        p = Path(path) if Path(path).is_absolute() else self.base / path
        if not p.is_file():
            out({"status": "error", "error": f"image not found: {path}"}, 2)
        try:
            from PIL import Image
            iw, ih = Image.open(p).size
        except Exception:  # noqa: BLE001
            iw, ih = 4, 3
        scale = min(w / iw, h / ih)
        dw, dh = iw * scale, ih * scale
        s.shapes.add_picture(str(p), Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2), Inches(dw), Inches(dh))

    # -- layouts
    def build(self, sl):
        kind = sl.get("layout") or ("title" if self.n == 0 else "bullets")
        s = self.slide()
        getattr(self, "l_" + kind.replace("-", "_"))(s, sl)
        if kind not in ("title", "section", "closing"):
            self.footer(s)
        self.notes(s, sl)

    def l_title(self, s, sl):
        self.bar(s, 0, 0, W_IN, H_IN, self.accent)
        self.bar(s, 0.9, 4.15, 1.6, 0.08, self.accent2)
        self.text(s, 0.9, 1.9, 11.5, 2.2, sl.get("title", ""), size=42, bold=True, color="FFFFFF", anchor="bottom")
        if sl.get("subtitle"):
            self.text(s, 0.9, 4.4, 11.5, 1.0, sl["subtitle"], size=20, color="E6E6E6")
        if sl.get("byline"):
            self.text(s, 0.9, 6.3, 11.5, 0.5, sl["byline"], size=14, color="D0D0D0")

    def l_section(self, s, sl):
        self.bar(s, 0, 0, 4.2, H_IN, self.accent)
        if sl.get("number") is not None:
            self.text(s, 0.6, 2.6, 3.2, 1.6, str(sl["number"]), size=72, bold=True, color="FFFFFF", align="left")
        self.text(s, 4.8, 2.5, 8, 1.4, sl.get("title", ""), size=36, bold=True, anchor="middle")
        if sl.get("subtitle"):
            self.text(s, 4.8, 3.9, 8, 1.0, sl["subtitle"], size=18, color=self.muted)

    def l_bullets(self, s, sl):
        self.title(s, sl.get("title", ""), sl.get("subtitle"))
        top = 1.9 if sl.get("subtitle") else 1.55
        self.text(s, 0.85, top, 11.8, 6.85 - top, sl.get("bullets", []), size=sl.get("size", 22), bullets=True)

    def l_two_column(self, s, sl):
        self.title(s, sl.get("title", ""), sl.get("subtitle"))
        for i, col in enumerate((sl.get("left", {}), sl.get("right", {}))):
            x = 0.85 + i * 6.1
            if col.get("heading"):
                self.text(s, x, 1.6, 5.7, 0.55, col["heading"], size=20, bold=True, color=self.accent)
                self.bar(s, x, 2.15, 5.6, 0.03, "D0D0D0")
            if col.get("image"):
                self.image(s, col["image"], x, 2.3, 5.6, 4.4)
            else:
                self.text(s, x, 2.3, 5.6, 4.5, col.get("bullets", []), size=col.get("size", 18), bullets=True)

    def l_image(self, s, sl):
        self.title(s, sl.get("title", ""))
        cap = sl.get("caption")
        self.image(s, sl["image"], 0.85, 1.55, 11.6, 4.9 if cap else 5.3)
        if cap:
            self.text(s, 0.85, 6.5, 11.6, 0.45, cap, size=14, color=self.muted, align="center")

    def l_chart(self, s, sl):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt
        self.title(s, sl.get("title", ""), sl.get("subtitle"))
        ch = sl["chart"]
        data = CategoryChartData()
        data.categories = ch["categories"]
        for name, vals in ch["series"].items():
            data.add_series(name, vals)
        kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
                "stacked": XL_CHART_TYPE.COLUMN_STACKED}[ch.get("type", "column")]
        top = 1.95 if sl.get("subtitle") else 1.6
        gf = s.shapes.add_chart(kind, Inches(0.85), Inches(top), Inches(11.6), Inches(6.75 - top), data)
        c = gf.chart
        c.font.size, c.font.name = Pt(14), self.font
        multi = len(ch["series"]) > 1 or ch.get("type") == "pie"
        c.has_legend = multi
        if multi:
            c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        palette = [self.accent, self.accent2, "7F7F7F", "2E75B6", "548235", "BF9000"]
        if ch.get("type") != "pie":
            for i, ser in enumerate(c.plots[0].series):
                fill = ser.format.line if ch.get("type") == "line" else ser.format.fill
                if ch.get("type") == "line":
                    fill.color.rgb = rgb(palette[i % len(palette)])
                    fill.width = Pt(2.5)
                else:
                    fill.solid()
                    fill.fore_color.rgb = rgb(palette[i % len(palette)])
            c.value_axis.has_major_gridlines = True
            c.value_axis.major_gridlines.format.line.color.rgb = rgb("E3E3E3")
            if ch.get("y_title"):
                c.value_axis.has_title = True
                c.value_axis.axis_title.text_frame.text = ch["y_title"]
            if ch.get("labels", True):
                pl = c.plots[0]
                pl.has_data_labels = True
                pl.data_labels.font.size = Pt(12)
                if ch.get("number_format"):
                    pl.data_labels.number_format = ch["number_format"]
                    pl.data_labels.number_format_is_linked = False
        else:
            pl = c.plots[0]
            pl.has_data_labels = True
            pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "0%", False
            pl.data_labels.show_percentage = True
            for i, pt in enumerate(pl.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
        if ch.get("source"):
            self.text(s, 0.85, 6.7, 11.6, 0.3, ch["source"], size=10, color=self.muted)

    def l_table(self, s, sl):
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches
        self.title(s, sl.get("title", ""), sl.get("subtitle"))
        t = sl["table"]
        header, rows = t.get("header") or [], t.get("rows") or []
        ncol = max([len(header)] + [len(r) for r in rows])
        nrow = len(rows) + (1 if header else 0)
        top = 1.95 if sl.get("subtitle") else 1.6
        h = min(0.5 * nrow, 6.75 - top)
        gf = s.shapes.add_table(nrow, ncol, Inches(0.85), Inches(top), Inches(11.6), Inches(h))
        tbl = gf.table
        size = 16 if nrow <= 8 else 13 if nrow <= 12 else 11
        all_rows = ([header] if header else []) + rows
        for i, r in enumerate(all_rows):
            for j in range(ncol):
                cell = tbl.cell(i, j)
                v = r[j] if j < len(r) else ""
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                head = bool(header) and i == 0
                self._runs(p, "" if v is None else str(v), size, head, "FFFFFF" if head else self.fg)
                if isinstance(v, (int, float)) or re.fullmatch(r"[-+]?[\d,.]+%?", str(v or "")):
                    p.alignment = PP_ALIGN.RIGHT
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.accent if head else ("F2F2F2" if i % 2 == 0 else "FFFFFF"))
                cell.margin_left = cell.margin_right = Inches(0.08)

    def l_quote(self, s, sl):
        self.bar(s, 0.9, 1.8, 0.12, 3.2, self.accent2)
        self.text(s, 1.3, 1.6, 10.8, 3.6, sl.get("quote", ""), size=30, anchor="middle")
        if sl.get("by"):
            self.text(s, 1.3, 5.4, 10.8, 0.6, "— " + sl["by"], size=16, color=self.muted)

    def l_big_number(self, s, sl):
        self.title(s, sl.get("title", ""))
        items = sl.get("numbers", [])[:4]
        w = 11.6 / max(len(items), 1)
        for i, it in enumerate(items):
            x = 0.85 + i * w
            self.text(s, x, 2.3, w - 0.3, 1.6, str(it.get("value", "")), size=54, bold=True, color=self.accent,
                      align="center", anchor="bottom")
            self.text(s, x, 4.0, w - 0.3, 1.4, it.get("label", ""), size=18, color=self.muted, align="center")

    def l_closing(self, s, sl):
        self.bar(s, 0, 0, W_IN, H_IN, self.accent)
        self.text(s, 0.9, 2.4, 11.5, 1.6, sl.get("title", "謝謝" if self.lang.startswith("zh") else "Thank you"),
                  size=44, bold=True, color="FFFFFF", anchor="middle", align="center")
        if sl.get("subtitle"):
            self.text(s, 0.9, 4.1, 11.5, 1.0, sl["subtitle"], size=18, color="E6E6E6", align="center")


def build(spec_path: Path, out_path: Path):
    need()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    deck = Deck(spec, spec_path.parent)
    for sl in spec.get("slides", []):
        layout = sl.get("layout") or ""
        if layout and not hasattr(deck, "l_" + layout.replace("-", "_")):
            out({"status": "error", "error": f"unknown layout {layout!r}; use title, section, bullets, two-column, "
                                             "image, chart, table, quote, big-number, closing"}, 2)
        deck.build(sl)
    if spec.get("title"):
        deck.prs.core_properties.title = spec["title"]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(str(out_path))
    out({"status": "ok", "file": str(out_path), "slides": deck.n, "east_asia_font": deck.east,
         "next": f"run.py check {out_path}  then  run.py render {out_path} <dir> and look at every slide"})


# ------------------------------------------------------------------ inspect / check

def _load(path: Path):
    need()
    from pptx import Presentation
    try:
        return Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        out({"status": "error", "error": f"cannot open {path.name} as .pptx: {type(e).__name__}: {e}"}, 2)


def _shape_title(slide) -> str:
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        return slide.shapes.title.text_frame.text.strip()
    # drawn decks: the text in the largest font
    best = ("", 0)
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            size = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
            if size > best[1]:
                best = (sh.text_frame.text.strip(), size)
    return best[0]


def inspect(path: Path):
    prs = _load(path)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        text = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        slides.append({"n": i, "title": _shape_title(s), "text": text[:12],
                       "pictures": sum(1 for sh in s.shapes if sh.shape_type == 13),
                       "charts": sum(1 for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart),
                       "tables": sum(1 for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table),
                       "notes": s.notes_slide.notes_text_frame.text.strip()[:400] if s.has_notes_slide else ""})
    out({"status": "ok", "size_in": [round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)],
         "slides": slides})


def _needed_height(tf, width_pt: float) -> float:
    total = 0.0
    for p in tf.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18.0)
        text = "".join(r.text for r in p.runs) or " "
        units = sum(1.0 if CJK.match(ch) else 0.55 for ch in text)
        indent = 0.0
        pPr = p._p.pPr
        if pPr is not None and pPr.get("marL"):
            indent = int(pPr.get("marL")) / 12700
        per_line = max((width_pt - indent) / size, 1)
        lines = max(1, math.ceil(units / per_line))
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
        before = p.space_before.pt if p.space_before is not None else 0
        total += lines * size * 1.2 * spacing + before
    return total


def check(path: Path):
    prs = _load(path)
    SW, SH = prs.slide_width, prs.slide_height
    problems, notes = [], []
    for i, s in enumerate(prs.slides, 1):
        words = 0
        if not _shape_title(s) and i > 1:
            notes.append(f"slide {i}: no title")
        for sh in s.shapes:
            if sh.left is not None and sh.width is not None:
                if sh.left < -9144 or sh.top < -9144 or sh.left + sh.width > SW + 9144 or sh.top + sh.height > SH + 9144:
                    problems.append(f"slide {i}: 「{getattr(sh, 'name', '')}」 sticks out of the slide")
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
                problems.append(f"slide {i}: empty placeholder 「{sh.name}」 (shows 'Click to add …' in edit view)")
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            text = tf.text
            words += len(re.findall(r"[A-Za-z]+", text)) + len(CJK.findall(text)) / 2
            if PLACEHOLDER.search(text):
                problems.append(f"slide {i}: placeholder text 「{text.strip()[:40]}」")
            sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
            if sizes and min(sizes) < 10 and text.strip():
                problems.append(f"slide {i}: text at {min(sizes):.0f} pt is too small to read when presented")
            if sh.width and sh.height and text.strip() and tf.word_wrap is not False:
                need_pt = _needed_height(tf, sh.width / 12700 - 7)
                have_pt = sh.height / 12700
                if need_pt > have_pt * 1.15 + 6:
                    problems.append(f"slide {i}: text overflows its box (needs ~{need_pt:.0f} pt, box {have_pt:.0f} pt): "
                                    f"「{text.strip()[:30]}…」 - shorten it, split the slide or use a smaller size")
            bullets = [p for p in tf.paragraphs if "".join(r.text for r in p.runs).strip()]
            if len(bullets) > 8:
                notes.append(f"slide {i}: {len(bullets)} lines in one box - consider splitting")
        if words > 120:
            notes.append(f"slide {i}: about {int(words)} words - slides read best under ~60")
    out({"status": "fix" if problems else "ok", "problems": problems[:30], "notes": notes[:30],
         "slides": len(prs.slides)}, 1 if problems else 0)


def render(path: Path, outdir: Path):
    outdir.mkdir(parents=True, exist_ok=True)
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    exe = shutil.which("soffice") or shutil.which("libreoffice") or (mac if Path(mac).exists() else None)
    if not exe:
        out({"status": "error", "error": "LibreOffice (soffice) is not installed: cannot draw the slides",
             "fix": "macOS: brew install --cask libreoffice; Linux: apt install libreoffice-impress"}, 2)
    subprocess.run([exe, "--headless", f"-env:UserInstallation=file://{outdir.resolve()}/.lo", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(path)], capture_output=True, timeout=240)
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.is_file():
        out({"status": "error", "error": "LibreOffice could not convert the deck"}, 2)
    pngs = []
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", "60", str(pdf), str(outdir / "slide")], capture_output=True,
                       timeout=240)
        pngs = sorted((str(p) for p in outdir.glob("slide*.png")), key=lambda p: int(re.sub(r"\D", "", Path(p).stem) or 0))
    out({"status": "ok", "pdf": str(pdf), "slides": pngs,
         "next": "read_file every slide image: overlaps, overflow, contrast, alignment, empty areas"})


def main(argv):
    if not argv:
        out({"status": "error", "error": __doc__}, 2)
    cmd, args = argv[0], argv[1:]
    try:
        if cmd == "build" and len(args) == 2:
            build(Path(args[0]), Path(args[1]))
        elif cmd == "inspect" and len(args) == 1:
            inspect(Path(args[0]))
        elif cmd == "check" and len(args) == 1:
            check(Path(args[0]))
        elif cmd == "render" and len(args) == 2:
            render(Path(args[0]), Path(args[1]))
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        out({"status": "error", "error": f"{type(e).__name__}: {e}"}, 2)
    out({"status": "error", "error": "usage:\n" + __doc__}, 2)


if __name__ == "__main__":
    main(sys.argv[1:])

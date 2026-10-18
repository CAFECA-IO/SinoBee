#!/usr/bin/env python3
"""Spreadsheets (.xlsx): build from a JSON description, read, recalculate, check.

    run.py build <book.json> <out.xlsx>     sheets with a styled header, typed values, live formulas, number
                                            formats, column widths, frozen header, filters, totals, native charts
    run.py inspect <file.xlsx> [--rows N]   sheets, size, formulas, merged cells, freeze panes, first rows
                                            (values AND formulas)
    run.py recalc <file.xlsx> <out.xlsx>    LibreOffice recalculates every formula and saves the values, so
                                            viewers that do not calculate (previews, read_file) show numbers
    run.py check <file.xlsx>                formula errors after recalculation (#DIV/0!, #REF!, #NAME? …),
                                            formulas pointing at empty cells, numbers stored as text, totals
                                            typed in as constants, sheets without a header

Every command prints JSON. Exit 0 = ok, 1 = check found problems, 2 = usage / input error.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ERRORS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
FORMATS = {"int": "#,##0", "number": "#,##0.00", "percent": "0.0%", "currency": "#,##0.00", "twd": '"NT$"#,##0',
           "usd": '"$"#,##0.00', "date": "yyyy-mm-dd", "text": "@"}
INPUT_FONT = "0000FF"        # convention: blue = typed-in inputs, black = formulas


def out(data: dict, code: int = 0):
    print(json.dumps(data, ensure_ascii=False, indent=1, default=str))
    sys.exit(code)


def need():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        out({"status": "error", "error": "openpyxl is not installed in this Python",
             "fix": f"{sys.executable} -m pip install openpyxl"}, 2)


def soffice():
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return shutil.which("soffice") or shutil.which("libreoffice") or (mac if Path(mac).exists() else None)


# ------------------------------------------------------------------ build

def _width(v) -> float:
    s = str(v if v is not None else "")
    return sum(2.0 if ord(ch) > 0x2E80 else 1.0 for ch in s)


def build(spec_path: Path, out_path: Path):
    need()
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    font_name = spec.get("font", "Calibri")
    made = []
    for sh in spec.get("sheets", []):
        ws = wb.create_sheet(str(sh.get("name", f"Sheet{len(wb.sheetnames) + 1}"))[:31])
        cols = sh.get("columns") or []          # [{"name", "format", "width", "input": bool}]
        header = [c["name"] if isinstance(c, dict) else str(c) for c in cols] or sh.get("header") or []
        r0 = 1
        if sh.get("title"):
            ws.cell(1, 1, sh["title"]).font = Font(name=font_name, bold=True, size=14)
            r0 = 3
        for j, h in enumerate(header, 1):
            c = ws.cell(r0, j, h)
            c.font = Font(name=font_name, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=sh.get("header_fill", "305496"))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=thin)
        rows = sh.get("rows") or []
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                col = cols[j - 1] if j - 1 < len(cols) and isinstance(cols[j - 1], dict) else {}
                fmt = col.get("format")
                if fmt:
                    c.number_format = FORMATS.get(fmt, fmt)
                is_formula = isinstance(v, str) and v.startswith("=")
                c.font = Font(name=font_name, color=INPUT_FONT if (col.get("input") and not is_formula) else None)
                c.border = Border(bottom=Side(style="hair", color="D9D9D9"))
        last = r0 + len(rows)
        if sh.get("totals") and rows:
            tr = last + 1
            ws.cell(tr, 1, sh.get("totals_label", "合計")).font = Font(name=font_name, bold=True)
            for col_name in sh["totals"]:
                j = header.index(col_name) + 1 if col_name in header else None
                if not j:
                    continue
                L = get_column_letter(j)
                c = ws.cell(tr, j, f"=SUM({L}{r0 + 1}:{L}{last})")
                c.font = Font(name=font_name, bold=True)
                c.border = Border(top=Side(style="thin"), bottom=Side(style="double"))
                fmt = cols[j - 1].get("format") if j - 1 < len(cols) and isinstance(cols[j - 1], dict) else None
                if fmt:
                    c.number_format = FORMATS.get(fmt, fmt)
        for j in range(1, max(len(header), max((len(r) for r in rows), default=0)) + 1):
            col = cols[j - 1] if j - 1 < len(cols) and isinstance(cols[j - 1], dict) else {}
            w = col.get("width") or min(max([_width(header[j - 1]) if j - 1 < len(header) else 0] +
                                            [_width(r[j - 1]) if j - 1 < len(r) and not str(r[j - 1]).startswith("=")
                                             else 10 for r in rows[:200]]) + 2, 60)
            ws.column_dimensions[get_column_letter(j)].width = max(w, 8)
        if header and sh.get("freeze", True):
            ws.freeze_panes = ws.cell(r0 + 1, 1)
        if header and rows and sh.get("filter", True):
            ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(header))}{last}"
        for ch in sh.get("charts", []):
            kind = ch.get("type", "bar")
            chart = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}[kind]()
            if kind == "bar":
                chart.type = "bar"
            if kind == "column":
                chart.type = "col"
            chart.title = ch.get("title")
            if kind != "pie":
                chart.y_axis.title = ch.get("y_title")
                chart.x_axis.title = ch.get("x_title")
            cat_col = header.index(ch["x"]) + 1
            for ycol in ch["y"] if isinstance(ch["y"], list) else [ch["y"]]:
                j = header.index(ycol) + 1
                chart.add_data(Reference(ws, min_col=j, min_row=r0, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=cat_col, min_row=r0 + 1, max_row=last))
            chart.width, chart.height = ch.get("width_cm", 18), ch.get("height_cm", 9)
            ws.add_chart(chart, ch.get("at", f"{get_column_letter(len(header) + 2)}{r0}"))
        made.append({"sheet": ws.title, "rows": len(rows), "header_row": r0})
    for name, ref in (spec.get("names") or {}).items():
        from openpyxl.workbook.defined_name import DefinedName
        dn = DefinedName(name, attr_text=ref)
        try:
            wb.defined_names[name] = dn
        except TypeError:
            wb.defined_names.append(dn)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    out({"status": "ok", "file": str(out_path), "sheets": made,
         "next": f"run.py check {out_path}  (formulas are stored without values until recalculated: "
                 f"run.py recalc {out_path} {out_path} if the reader needs the values)"})


# ------------------------------------------------------------------ recalc / inspect / check

def _recalc(src: Path, dst_dir: Path) -> Path | None:
    exe = soffice()
    if not exe:
        return None
    subprocess.run([exe, "--headless", f"-env:UserInstallation=file://{dst_dir.resolve()}/.lo", "--convert-to",
                    "xlsx", "--outdir", str(dst_dir), str(src)], capture_output=True, timeout=240)
    p = dst_dir / (src.stem + ".xlsx")
    return p if p.is_file() else None


def recalc(src: Path, dst: Path):
    with tempfile.TemporaryDirectory() as tmp:
        p = _recalc(src, Path(tmp))
        if not p:
            out({"status": "error", "error": "LibreOffice is not installed (or failed): cannot recalculate",
                 "fix": "macOS: brew install --cask libreoffice; Linux: apt install libreoffice-calc"}, 2)
        shutil.copy(p, dst)
    out({"status": "ok", "file": str(dst), "note": "values are now stored; check charts and conditional formats "
                                                    "survived the LibreOffice round trip"})


def inspect(path: Path, nrows: int = 15):
    need()
    import openpyxl
    wb = openpyxl.load_workbook(path)
    vals = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        wv = vals[ws.title]
        formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        first = []
        for r, (row, vrow) in enumerate(zip(ws.iter_rows(max_row=nrows), wv.iter_rows(max_row=nrows)), 1):
            cells = []
            for c, v in zip(row, vrow):
                if c.value is None:
                    continue
                cells.append(f"{c.coordinate}={c.value!r}" + (f"→{v.value!r}" if c.value != v.value else ""))
            if cells:
                first.append("  ".join(cells)[:600])
        sheets.append({"name": ws.title, "size": ws.dimensions, "max_row": ws.max_row, "max_column": ws.max_column,
                       "formulas": formulas, "merged": [str(m) for m in ws.merged_cells.ranges][:20],
                       "freeze_panes": ws.freeze_panes, "charts": len(getattr(ws, "_charts", [])),
                       "first_rows": first})
    out({"status": "ok", "sheets": sheets, "defined_names": list(wb.defined_names.keys())
         if hasattr(wb.defined_names, "keys") else []})


def check(path: Path):
    need()
    import openpyxl
    problems, notes = [], []
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:  # noqa: BLE001
        out({"status": "fix", "problems": [f"cannot open {path.name}: {type(e).__name__}: {e}"]}, 1)
    with tempfile.TemporaryDirectory() as tmp:
        calc = _recalc(path, Path(tmp))
        values = openpyxl.load_workbook(calc or path, data_only=True)
        if not calc:
            notes.append("LibreOffice not installed: formulas were not recalculated, only cached values checked")
        for ws in wb.worksheets:
            wv = values[ws.title] if ws.title in values.sheetnames else None
            errors, text_nums, empty_refs = [], [], []
            numeric_cols: dict[int, int] = {}
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    cv = wv[c.coordinate].value if wv is not None else None
                    if isinstance(cv, str) and cv.strip() in ERRORS:
                        errors.append(f"{c.coordinate} {cv} ({v})")
                    elif isinstance(v, str) and v.startswith("=") and "#REF!" in v:
                        errors.append(f"{c.coordinate} #REF! in formula {v}")
                    if isinstance(v, (int, float)):
                        numeric_cols[c.column] = numeric_cols.get(c.column, 0) + 1
                    if isinstance(v, str) and re.fullmatch(r"\s*[-+]?\d[\d,]*(\.\d+)?\s*", v) and c.row > 1:
                        text_nums.append(c.coordinate)
                    if isinstance(v, str) and v.startswith("=") and wv is not None:
                        for ref in re.findall(r"(?<![A-Z!:$])\$?([A-Z]{1,3})\$?(\d{1,7})(?![\d(:])", v.split("!")[-1]):
                            coord = f"{ref[0]}{ref[1]}"
                            try:
                                if ws[coord].value is None:
                                    empty_refs.append(f"{c.coordinate}→{coord}")
                            except ValueError:
                                pass
            if errors:
                problems.append(f"{ws.title}: {len(errors)} formula error(s): " + "; ".join(errors[:8]))
            if text_nums:
                problems.append(f"{ws.title}: numbers stored as text in {', '.join(text_nums[:8])}"
                                + (f" (+{len(text_nums) - 8})" if len(text_nums) > 8 else "")
                                + " - write them as numbers so formulas and sorting work")
            if empty_refs:
                notes.append(f"{ws.title}: formulas refer to empty cells: {', '.join(empty_refs[:8])}")
            header = [c.value for c in next(ws.iter_rows(max_row=1), [])]
            if ws.max_row > 3 and not any(isinstance(h, str) and h.strip() for h in header) and not ws["A1"].value:
                notes.append(f"{ws.title}: no header row")
            # totals typed in as constants: a row labelled total/合計 with numbers but no formulas
            for row in ws.iter_rows():
                label = next((str(c.value) for c in row if isinstance(c.value, str)), "")
                if re.search(r"(?i)^(total|sum|合計|總計|小計)$", label.strip()):
                    nums = [c for c in row if isinstance(c.value, (int, float))]
                    if nums and not any(isinstance(c.value, str) and c.value.startswith("=") for c in row):
                        problems.append(f"{ws.title}: the {label} row {row[0].row} has typed-in numbers; use =SUM() so "
                                        "it stays right when inputs change")
    out({"status": "fix" if problems else "ok", "problems": problems, "notes": notes,
         "recalculated": bool(calc)}, 1 if problems else 0)


def main(argv):
    if not argv:
        out({"status": "error", "error": __doc__}, 2)
    cmd, args = argv[0], argv[1:]
    try:
        if cmd == "build" and len(args) == 2:
            build(Path(args[0]), Path(args[1]))
        elif cmd == "inspect" and args:
            inspect(Path(args[0]), int(args[args.index("--rows") + 1]) if "--rows" in args else 15)
        elif cmd == "recalc" and len(args) == 2:
            recalc(Path(args[0]), Path(args[1]))
        elif cmd == "check" and len(args) == 1:
            check(Path(args[0]))
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        out({"status": "error", "error": f"{type(e).__name__}: {e}"}, 2)
    out({"status": "error", "error": "usage:\n" + __doc__}, 2)


if __name__ == "__main__":
    main(sys.argv[1:])

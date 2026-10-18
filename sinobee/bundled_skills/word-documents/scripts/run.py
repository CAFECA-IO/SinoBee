#!/usr/bin/env python3
"""Word documents (.docx): build from a JSON outline, read, edit with tracked changes, check, render.

    run.py build <doc.json> <out.docx>      outline -> .docx (headings, paragraphs with **bold** *italic* `code`,
                                            lists, tables, images with captions, quotes, code, page breaks,
                                            table of contents, header / footer / page numbers, CJK fonts)
    run.py inspect <file.docx>              structure: headings, paragraph / table / image counts, styles, fonts,
                                            comments, tracked changes, and the text as Markdown
    run.py replace <in> <out> <old> <new> [--track] [--author NAME]
                                            replace text inside paragraphs (also across runs, keeping the first
                                            run's formatting); --track records it as a tracked change
    run.py check <file.docx>                problems a reader would hit (JSON {"status", "problems"})
    run.py render <file.docx> <outdir>      PNG of each page via LibreOffice + pdftoppm, for looking at it

Every command prints JSON. Exit code 0 on success, 1 when check finds problems, 2 on usage / input errors.
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

CJK = re.compile(r"[぀-ヿ㐀-鿿가-힯豈-﫿]")
PLACEHOLDER = re.compile(r"lorem ipsum|\bTODO\b|\bTBD\b|\bXXX\b|\[(?:insert|placeholder|your [a-z ]+)\]|替換成|請填入|待補",
                         re.I)
EAST_ASIA_DEFAULT = {"zh-Hant": "Microsoft JhengHei", "zh-Hans": "Microsoft YaHei", "ja": "Yu Gothic",
                     "ko": "Malgun Gothic"}


def out(data: dict, code: int = 0):
    print(json.dumps(data, ensure_ascii=False, indent=1))
    sys.exit(code)


def need():
    try:
        import docx  # noqa: F401
    except ImportError:
        out({"status": "error", "error": "python-docx is not installed in this Python",
             "fix": f"{sys.executable} -m pip install python-docx"}, 2)


# ------------------------------------------------------------------ build

def _runs(par, text: str, base_bold=False):
    """**bold**, *italic*, `code`, [text](url) inside a paragraph."""
    from docx.shared import Pt
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\[([^\]]+)\]\((https?://[^)]+)\)", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()]).bold = base_bold or None
        if m.group(1):
            par.add_run(m.group(1)).bold = True
        elif m.group(2):
            par.add_run(m.group(2)).italic = True
        elif m.group(3):
            r = par.add_run(m.group(3))
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            _hyperlink(par, m.group(4), m.group(5))
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:]).bold = base_bold or None


def _hyperlink(par, text, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    st = OxmlElement("w:rStyle")
    st.set(qn("w:val"), "Hyperlink")
    rpr.append(st)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "0563C1")
    rpr.append(col)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def _set_fonts(doc, latin: str, east: str, size: float):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title", "List Bullet",
                 "List Number", "Caption", "Quote", "Table Grid"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = latin
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        for k in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(k), latin)
        fonts.set(qn("w:eastAsia"), east)
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(k), None)
        if name == "Normal":
            st.font.size = Pt(size)


def _field(par, instr: str, placeholder: str = ""):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = par.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        elif kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


def _shade(cell, hex_color: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def build(spec_path: Path, out_path: Path):
    need()
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt, RGBColor
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    base = spec_path.parent
    doc = Document()
    page = spec.get("page", {})
    sec = doc.sections[0]
    if str(page.get("size", "A4")).upper() == "LETTER":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if page.get("orientation") == "landscape":
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    m = Cm(float(page.get("margins_cm", 2.5)))
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
    text_all = json.dumps(spec, ensure_ascii=False)
    fonts = spec.get("fonts", {})
    lang = spec.get("lang") or ("zh-Hant" if CJK.search(text_all) else "en")
    east = fonts.get("east_asia") or EAST_ASIA_DEFAULT.get(lang, "Microsoft JhengHei")
    _set_fonts(doc, fonts.get("latin", "Calibri"), east, float(fonts.get("size", 11)))
    if spec.get("title"):
        doc.core_properties.title = spec["title"]
    if spec.get("author"):
        doc.core_properties.author = spec["author"]
    if spec.get("header"):
        sec.header.paragraphs[0].text = spec["header"]
    if spec.get("footer") or spec.get("page_numbers", True):
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if spec.get("footer"):
            fp.add_run(spec["footer"] + "    ")
        if spec.get("page_numbers", True):
            _field(fp, "PAGE", "1")
    if spec.get("title"):
        doc.add_heading(spec["title"], 0)
    if spec.get("subtitle"):
        p = doc.add_paragraph()
        r = p.add_run(spec["subtitle"])
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if spec.get("toc"):
        doc.add_paragraph().add_run("目錄" if lang.startswith("zh") else "Contents").bold = True
        _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "（開啟文件後按 F9 更新目錄）" if lang.startswith("zh")
               else "(Right-click > Update Field to build the table of contents)")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    fig = tab = 0
    for b in spec.get("blocks", []):
        if not isinstance(b, dict):
            b = {"p": str(b)}
        for lvl in (1, 2, 3, 4):
            if f"h{lvl}" in b:
                doc.add_heading(b[f"h{lvl}"], lvl)
        if "p" in b:
            par = doc.add_paragraph()
            _runs(par, str(b["p"]))
            if b.get("align") == "center":
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "bullets" in b or "numbered" in b:
            items = b.get("bullets") or b.get("numbered")
            style = "List Bullet" if "bullets" in b else "List Number"
            for it in items:
                level = 0
                if isinstance(it, dict):
                    level, it = int(it.get("level", 0)), it.get("text", "")
                st = style if level == 0 else f"{style} {min(level + 1, 3)}"
                try:
                    par = doc.add_paragraph(style=st)
                except KeyError:
                    par = doc.add_paragraph(style=style)
                _runs(par, str(it))
        if "quote" in b:
            par = doc.add_paragraph(style="Quote") if "Quote" in [s.name for s in doc.styles] else doc.add_paragraph()
            _runs(par, str(b["quote"]))
        if "code" in b:
            for line in str(b["code"]).splitlines() or [""]:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                r = par.add_run(line or " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9)
        if "table" in b:
            t = b["table"]
            header, rows = t.get("header") or [], t.get("rows") or []
            ncol = max([len(header)] + [len(r) for r in rows] or [1])
            if t.get("caption"):
                tab += 1
                cap = doc.add_paragraph(style="Caption")
                cap.add_run(f"{'表' if lang.startswith('zh') else 'Table'} {tab}  {t['caption']}")
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if header:
                cells = table.add_row().cells
                for i, h in enumerate(header):
                    cells[i].text = ""
                    _runs(cells[i].paragraphs[0], str(h), base_bold=True)
                    _shade(cells[i], t.get("header_fill", "D9E2F3"))
                _repeat_header(table.rows[0])
            for r in rows:
                cells = table.add_row().cells
                for i, v in enumerate(r[:ncol]):
                    cells[i].text = ""
                    _runs(cells[i].paragraphs[0], "" if v is None else str(v))
                    if isinstance(v, (int, float)) or re.fullmatch(r"[-+]?[\d,.]+%?", str(v or "")):
                        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for i, w in enumerate(t.get("widths_cm") or []):
                for row in table.rows:
                    if i < len(row.cells):
                        row.cells[i].width = Cm(float(w))
            doc.add_paragraph()
        if "image" in b:
            im = b["image"] if isinstance(b["image"], dict) else {"path": b["image"]}
            p = (base / im["path"]) if not Path(im["path"]).is_absolute() else Path(im["path"])
            if not p.is_file():
                out({"status": "error", "error": f"image not found: {im['path']}"}, 2)
            doc.add_picture(str(p), width=Cm(float(im.get("width_cm", 15))))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if im.get("caption"):
                fig += 1
                cap = doc.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run(f"{'圖' if lang.startswith('zh') else 'Figure'} {fig}  {im['caption']}")
        if b.get("page_break"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    out({"status": "ok", "file": str(out_path), "east_asia_font": east, "figures": fig, "tables": tab,
         "next": f"run.py check {out_path}  then  run.py render {out_path} <dir> and look at the pages"})


# ------------------------------------------------------------------ inspect / check

def _load(path: Path):
    need()
    from docx import Document
    try:
        return Document(str(path))
    except Exception as e:  # noqa: BLE001
        out({"status": "error", "error": f"cannot open {path.name} as .docx: {type(e).__name__}: {e}"}, 2)


def _text(path: Path) -> str:
    """Markdown of the document (same reader SinoBee's read_file uses, when available)."""
    try:
        from sinobee.docread import extract
        return extract(path)
    except Exception:  # noqa: BLE001 - running outside SinoBee
        doc = _load(path)
        return "\n".join(p.text for p in doc.paragraphs)


def facts(path: Path) -> dict:
    from docx.oxml.ns import qn
    doc = _load(path)
    body = doc.element.body
    heads = []
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else ""
        m = re.match(r"Heading (\d)|Title", name)
        if m and p.text.strip():
            heads.append((int(m.group(1) or 0), p.text.strip()))
    fonts = set()
    east = set()
    for st in doc.styles:
        rpr = getattr(st.element, "rPr", None)
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rf is not None:
            fonts.update(v for k, v in rf.attrib.items() if not k.endswith("Theme"))
            if rf.get(qn("w:eastAsia")):
                east.add(rf.get(qn("w:eastAsia")))
    for r in body.iter(qn("w:rFonts")):
        if r.get(qn("w:eastAsia")):
            east.add(r.get(qn("w:eastAsia")))
    blank_run, longest_blank = 0, 0
    for p in doc.paragraphs:
        blank_run = blank_run + 1 if not p.text.strip() and not p._p.xpath(".//w:drawing|.//w:br") else 0
        longest_blank = max(longest_blank, blank_run)
    return {
        "doc": doc, "headings": heads, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
        "images": len(body.xpath(".//w:drawing")), "east_asia_fonts": sorted(east), "fonts": sorted(fonts),
        "comments": _count(path, "word/comments.xml", "w:comment"),
        "insertions": len(body.xpath(".//w:ins")), "deletions": len(body.xpath(".//w:del")),
        "longest_blank_run": longest_blank, "sections": len(doc.sections),
    }


def _count(path: Path, part: str, tag: str) -> int:
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            return z.read(part).decode("utf-8", "replace").count(f"<{tag} ")
    except KeyError:
        return 0


def inspect(path: Path):
    f = facts(path)
    f.pop("doc")
    text = _text(path)
    out({"status": "ok", **f, "headings": [f"{'#' * max(lvl, 1)} {t}" for lvl, t in f["headings"]][:80],
         "characters": len(text), "text": text[:12000] + ("\n… (truncated; read_file shows all)" if len(text) > 12000 else "")})


def check(path: Path):
    f = facts(path)
    doc = f["doc"]
    problems = []
    text = _text(path)
    if len(re.sub(r"\s", "", text)) < 20:
        problems.append("the document is (almost) empty")
    prev = None
    for lvl, t in f["headings"]:
        if prev is not None and lvl > prev + 1 and lvl > 1:
            problems.append(f"heading level jumps from {prev} to {lvl} at「{t[:40]}」: use the next level down")
        prev = lvl
    for m in PLACEHOLDER.finditer(text):
        problems.append(f"placeholder text left in the document: 「{text[max(0, m.start() - 20):m.end() + 20].strip()}」")
        break
    if CJK.search(text) and not f["east_asia_fonts"]:
        problems.append("CJK text but no East Asian font is set (w:eastAsia): Word falls back to a font that may "
                        "lack the characters or mix styles - build with run.py build, or set the eastAsia font on "
                        "the Normal and Heading styles")
    if f["longest_blank_run"] >= 3:
        problems.append(f"{f['longest_blank_run']} empty paragraphs in a row: use spacing or a page break instead")
    for i, t in enumerate(doc.tables, 1):
        if len(t.rows) and all(not c.text.strip() for c in t.rows[0].cells):
            problems.append(f"table {i} has an empty first row (no header)")
    out({"status": "fix" if problems else "ok", "problems": problems,
         "summary": {k: v for k, v in f.items() if k != "doc" and k != "headings"},
         "headings": len(f["headings"])}, 1 if problems else 0)


# ------------------------------------------------------------------ replace (optionally tracked)

def replace(src: Path, dst: Path, old: str, new: str, track: bool, author: str):
    need()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = _load(src)
    now = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    rid = [1000]
    count = 0

    def paragraphs(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    def mark(tag, run_el):
        el = OxmlElement(tag)
        rid[0] += 1
        el.set(qn("w:id"), str(rid[0]))
        el.set(qn("w:author"), author)
        el.set(qn("w:date"), now)
        el.append(run_el)
        return el

    containers = [doc] + [s.header for s in doc.sections] + [s.footer for s in doc.sections]
    for c in containers:
        for p in paragraphs(c):
            while old in p.text:
                runs = p.runs
                full = "".join(r.text for r in runs)
                start = full.find(old)
                if start < 0:
                    break
                end = start + len(old)
                # split runs so [start, end) is exactly a run sequence
                pos, first, last = 0, None, None
                for r in runs:
                    n = len(r.text)
                    if first is None and pos + n > start:
                        first = r
                    if pos < end <= pos + n:
                        last = r
                    pos += n
                if first is None or last is None:
                    break
                # rebuild: text before stays in `first`, after goes to a copy of `last`
                pos = 0
                seq, inside = [], False
                for r in runs:
                    if r is first:
                        inside = True
                    if inside:
                        seq.append((r, pos))
                    pos += len(r.text)
                    if r is last:
                        break
                f_run, f_pos = seq[0]
                l_run, l_pos = seq[-1]
                before = f_run.text[:start - f_pos]
                after = l_run.text[end - l_pos:]
                old_texts = []
                for r, rp in seq:
                    s0 = max(start - rp, 0)
                    e0 = min(end - rp, len(r.text))
                    old_texts.append(r.text[s0:e0])
                tail = copy.deepcopy(l_run._r)
                for r, _ in seq[1:]:
                    r._r.getparent().remove(r._r)
                f_run.text = before
                anchor = f_run._r
                if track:
                    for t in old_texts:
                        d = copy.deepcopy(f_run._r)
                        for x in d.findall(qn("w:t")):
                            d.remove(x)
                        dt_el = OxmlElement("w:delText")
                        dt_el.set(qn("xml:space"), "preserve")
                        dt_el.text = t
                        d.append(dt_el)
                        wrapped = mark("w:del", d)
                        anchor.addnext(wrapped)
                        anchor = wrapped
                if new:
                    n_run = copy.deepcopy(f_run._r)
                    for x in n_run.findall(qn("w:t")):
                        n_run.remove(x)
                    t_el = OxmlElement("w:t")
                    t_el.set(qn("xml:space"), "preserve")
                    t_el.text = new
                    n_run.append(t_el)
                    node = mark("w:ins", n_run) if track else n_run
                    anchor.addnext(node)
                    anchor = node
                for x in tail.findall(qn("w:t")):
                    tail.remove(x)
                t_el = OxmlElement("w:t")
                t_el.set(qn("xml:space"), "preserve")
                t_el.text = after
                tail.append(t_el)
                anchor.addnext(tail)
                count += 1
                if track:
                    break        # the deleted text still reads as `old` in p.text only without delText; stop here
    doc.save(str(dst))
    out({"status": "ok" if count else "not_found", "replaced": count, "tracked": track, "file": str(dst)},
        0 if count else 1)


# ------------------------------------------------------------------ render

def render(path: Path, outdir: Path):
    outdir.mkdir(parents=True, exist_ok=True)
    exe = shutil.which("soffice") or shutil.which("libreoffice") or \
        ("/Applications/LibreOffice.app/Contents/MacOS/soffice"
         if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists() else None)
    if not exe:
        out({"status": "error", "error": "LibreOffice (soffice) is not installed: cannot draw the pages",
             "fix": "macOS: brew install --cask libreoffice; Linux: apt install libreoffice-writer"}, 2)
    subprocess.run([exe, "--headless", f"-env:UserInstallation=file://{outdir.resolve()}/.lo", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(path)], capture_output=True, timeout=180)
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.is_file():
        out({"status": "error", "error": "LibreOffice could not convert the document"}, 2)
    pngs = []
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), str(outdir / "page")], capture_output=True, timeout=180)
        pngs = sorted(str(p) for p in outdir.glob("page*.png"))
    out({"status": "ok", "pdf": str(pdf), "pages": pngs,
         "next": "read_file each page image and look at it: layout, fonts, tables, figures"})


def main(argv):
    if not argv:
        out({"status": "error", "error": __doc__}, 2)
    cmd, args = argv[0], argv[1:]
    try:
        if cmd == "build" and len(args) == 2:
            build(Path(args[0]), Path(args[1]))
        elif cmd == "inspect" and len(args) == 1:
            inspect(Path(args[0]))
        elif cmd == "check" and len(args) == 1:
            check(Path(args[0]))
        elif cmd == "render" and len(args) == 2:
            render(Path(args[0]), Path(args[1]))
        elif cmd == "replace" and len(args) >= 4:
            track = "--track" in args
            author = args[args.index("--author") + 1] if "--author" in args else "SinoBee"
            pos = [a for i, a in enumerate(args) if a != "--track" and a != "--author"
                   and (i == 0 or args[i - 1] != "--author")]
            replace(Path(pos[0]), Path(pos[1]), pos[2], pos[3], track, author)
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        out({"status": "error", "error": f"{type(e).__name__}: {e}"}, 2)
    out({"status": "error", "error": "usage:\n" + __doc__}, 2)


if __name__ == "__main__":
    main(sys.argv[1:])

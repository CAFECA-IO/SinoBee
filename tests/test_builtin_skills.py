"""The skills a general agent needs, filled in from a comparison with Claude's built-in skills: Word, Excel,
PowerPoint and PDF files, charts, web pages and code review - plus the core pieces they rely on: read_file reads
Office/PDF text instead of zip bytes, Python skill scripts run on SinoBee's own interpreter, a skill names the
packages it needs, and the visual check looks at produced documents' pages."""
import asyncio
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from sinobee.skills import spec
from sinobee.tools import REGISTRY, ToolContext
from sinobee.types import ToolCall

SK = Path(__file__).resolve().parent.parent / "sinobee" / "bundled_skills"
NEW = ["word-documents", "spreadsheets", "presentations", "pdf-documents", "data-visualization",
       "frontend-design", "code-review"]
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def run(skill, *args, cwd=None):
    r = subprocess.run([sys.executable, str(SK / skill / "scripts" / "run.py"), *map(str, args)],
                       capture_output=True, text=True, cwd=cwd, timeout=300)
    return r.returncode, json.loads(r.stdout)


def need(mod):
    return pytest.importorskip(mod)


# ---------- the skills themselves ----------

@pytest.mark.parametrize("name", NEW)
def test_new_skills_are_valid_and_complete(name):
    p = spec.parse((SK / name / "SKILL.md").read_text(encoding="utf-8"))
    assert not p.diagnostics and not spec.validate(p.frontmatter, name)
    assert not spec.quality_warnings(p.frontmatter, p.body)
    desc = p.frontmatter["description"]
    assert desc.rstrip().endswith((")", ".")), f"description cut short (a ' #' starts a YAML comment): {desc[-40:]}"
    meta = p.frontmatter["metadata"]
    for key in ("sinobee-checks",):
        if key in meta:
            assert (SK / name / meta[key].split()[0]).is_file()
    assert (SK / name / meta.get("sinobee-entry", "scripts/run.py")).is_file()


def test_every_bundled_description_survives_yaml():
    for d in SK.iterdir():
        if (d / "SKILL.md").exists():
            p = spec.parse((d / "SKILL.md").read_text(encoding="utf-8"))
            raw = (d / "SKILL.md").read_text(encoding="utf-8").split("description:", 1)[1].split("\n", 1)[0].strip()
            if raw.startswith(("'", '"')):          # quoted: YAML comments cannot cut it
                continue
            assert len(p.frontmatter["description"]) == len(raw), f"{d.name}: ' #' in the description starts a comment"


# ---------- core: skill runtime ----------

def test_missing_python_names_pip_packages():
    from sinobee.tools.skill_tools import missing_python
    assert missing_python("json, docx_not_here=python-docx-x, also_missing") == ["python-docx-x", "also_missing"]
    assert missing_python(["json"]) == [] and missing_python(None) == []


def test_skill_scripts_run_on_sinobees_python(config):
    from sinobee.agent import Runtime
    rt = Runtime(config)
    ctx = ToolContext(config=config, skills=rt.skills)
    r = asyncio.run(REGISTRY.dispatch(ToolCall("1", "skill_run", {"name": "data-visualization",
                                                                  "args": "palette categorical 3"}), ctx))
    if "needs Python packages" in r.content:
        assert "pip install" in r.content and sys.executable in r.content
        return
    data = json.loads(r.content[r.content.index("{"):])
    assert data["colors"] == ["#0072B2", "#D55E00", "#009E73"]


# ---------- core: reading Office / PDF ----------

def _docx(tmp_path) -> Path:
    need("docx")
    (tmp_path / "doc.json").write_text(json.dumps({
        "title": "測試報告", "blocks": [{"h1": "摘要"}, {"p": "金額為一百萬元，**重要**。"},
                                    {"bullets": ["甲", "乙"]},
                                    {"table": {"header": ["項目", "數值"], "rows": [["功率", "2 MWt"]]}}]},
        ensure_ascii=False), encoding="utf-8")
    code, res = run("word-documents", "build", tmp_path / "doc.json", tmp_path / "r.docx")
    assert code == 0, res
    return tmp_path / "r.docx"


def test_read_file_returns_document_text_not_zip_bytes(config, tmp_path):
    from sinobee.docread import extract
    docx = _docx(tmp_path)
    text = extract(docx)
    assert "# 摘要" in text and "金額為一百萬元" in text and "| 功率 | 2 MWt |" in text and "- 甲" in text
    ws = Path(config.workspace)
    ws.mkdir(parents=True, exist_ok=True)
    shutil.copy(docx, ws / "r.docx")
    r = asyncio.run(REGISTRY.dispatch(ToolCall("1", "read_file", {"path": "r.docx"}), ToolContext(config=config)))
    assert "金額為一百萬元" in r.content and "PK\x03\x04" not in r.content


def test_xlsx_and_pptx_text(tmp_path):
    need("openpyxl")
    import openpyxl
    from sinobee.docread import extract
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "預算"
    ws.append(["項目", "金額"])
    ws.append(["伺服器", 100])
    ws["B3"] = "=SUM(B2:B2)"
    wb.save(tmp_path / "b.xlsx")
    text = extract(tmp_path / "b.xlsx")
    assert "## Sheet: 預算" in text and "A2: 伺服器" in text and "B3: =SUM(B2:B2) (not calculated)" in text
    need("pptx")
    (tmp_path / "d.json").write_text(json.dumps({"slides": [
        {"layout": "title", "title": "釷基熔鹽"},
        {"layout": "bullets", "title": "常壓運轉", "bullets": ["沸點高"], "notes": "強調安全"},
        {"layout": "table", "title": "參數", "table": {"header": ["項目", "值"], "rows": [["功率", "2"]]}}]},
        ensure_ascii=False), encoding="utf-8")
    assert run("presentations", "build", tmp_path / "d.json", tmp_path / "d.pptx")[0] == 0
    text = extract(tmp_path / "d.pptx")
    assert "## Slide 1" in text and "常壓運轉" in text and "Notes: 強調安全" in text and "| 功率 | 2 |" in text


def test_a_broken_office_file_is_explained(tmp_path):
    from sinobee.docread import extract
    (tmp_path / "x.docx").write_bytes(b"not a zip")
    assert "not a valid Office file" in extract(tmp_path / "x.docx")


# ---------- word ----------

def test_word_build_check_and_tracked_replace(tmp_path):
    docx = _docx(tmp_path)
    code, res = run("word-documents", "check", docx)
    assert code == 0 and res["status"] == "ok" and res["summary"]["east_asia_fonts"] == ["Microsoft JhengHei"]
    code, res = run("word-documents", "replace", docx, tmp_path / "t.docx", "一百萬元", "兩百萬元", "--track")
    assert code == 0 and res["replaced"] == 1
    from sinobee.docread import extract
    assert "[-一百萬元-]{+兩百萬元+}" in extract(tmp_path / "t.docx")
    # a hand-made document with the usual mistakes
    import docx as d
    doc = d.Document()
    doc.add_heading("報告", 1)
    doc.add_heading("細節", 3)
    doc.add_paragraph("內容待補，請填入數字 TODO")
    for _ in range(4):
        doc.add_paragraph("")
    doc.save(tmp_path / "bad.docx")
    code, res = run("word-documents", "check", tmp_path / "bad.docx")
    text = "\n".join(res["problems"])
    assert code == 1 and "jumps from 1 to 3" in text and "placeholder" in text and "empty paragraphs" in text


# ---------- spreadsheets ----------

def test_spreadsheet_build_and_typed_totals(tmp_path):
    need("openpyxl")
    (tmp_path / "b.json").write_text(json.dumps({"sheets": [{"name": "預算", "columns": [
        {"name": "項目"}, {"name": "單價", "format": "twd", "input": True}, {"name": "數量", "format": "int"},
        {"name": "小計", "format": "twd"}], "rows": [["A", 10, 2, "=B2*C2"], ["B", 5, 3, "=B3*C3"]],
        "totals": ["小計"], "charts": [{"type": "column", "title": "小計", "x": "項目", "y": "小計"}]}]},
        ensure_ascii=False), encoding="utf-8")
    code, res = run("spreadsheets", "build", tmp_path / "b.json", tmp_path / "b.xlsx")
    assert code == 0, res
    code, res = run("spreadsheets", "check", tmp_path / "b.xlsx")
    assert code == 0, res
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["項目", "金額"])
    ws.append(["A", "1,200"])
    ws.append(["合計", 1200])
    wb.save(tmp_path / "bad.xlsx")
    text = "\n".join(run("spreadsheets", "check", tmp_path / "bad.xlsx")[1]["problems"])
    assert "numbers stored as text in B2" in text and "typed-in numbers" in text


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_spreadsheet_formula_errors_are_found_after_recalculation(tmp_path):
    need("openpyxl")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", "b", "a/b"])
    ws.append([1, 0, "=A2/B2"])
    ws["D1"] = "=NOSUCHFUNC(1)"
    wb.save(tmp_path / "e.xlsx")
    code, res = run("spreadsheets", "check", tmp_path / "e.xlsx")
    text = "\n".join(res["problems"])
    assert code == 1 and res["recalculated"] and "#DIV/0!" in text and "#NAME?" in text


# ---------- presentations ----------

def test_presentation_overflow_is_caught(tmp_path):
    need("pptx")
    long = "很長的文字" * 60
    (tmp_path / "d.json").write_text(json.dumps({"slides": [
        {"layout": "title", "title": "標題"},
        {"layout": "bullets", "title": "正常", "bullets": ["一", "二"]},
        {"layout": "bullets", "title": "太多字", "bullets": [long, long, long]}]}, ensure_ascii=False), encoding="utf-8")
    assert run("presentations", "build", tmp_path / "d.json", tmp_path / "d.pptx")[0] == 0
    code, res = run("presentations", "check", tmp_path / "d.pptx")
    assert code == 1 and len(res["problems"]) == 1 and res["problems"][0].startswith("slide 3: text overflows")
    (tmp_path / "e.json").write_text(json.dumps({"slides": [{"layout": "nope"}]}), encoding="utf-8")
    code, res = run("presentations", "build", tmp_path / "e.json", tmp_path / "e.pptx")
    assert code == 2 and "unknown layout" in res["error"]


# ---------- pdf ----------

def test_pdf_merge_split_fill_and_check(tmp_path):
    need("pypdf")
    rl = pytest.importorskip("reportlab.pdfgen.canvas")
    c = rl.Canvas(str(tmp_path / "form.pdf"))
    c.drawString(50, 750, "Name")
    c.acroForm.textfield(name="name", x=100, y=740, width=200, height=20)
    c.save()
    c = rl.Canvas(str(tmp_path / "two.pdf"))
    c.drawString(50, 750, "page one")
    c.showPage()
    c.drawString(50, 750, "page two")
    c.save()
    code, res = run("pdf-documents", "merge", tmp_path / "m.pdf", tmp_path / "form.pdf", tmp_path / "two.pdf")
    assert code == 0 and res["pages"] == 3
    code, res = run("pdf-documents", "split", tmp_path / "m.pdf", tmp_path / "parts", "--ranges", "1,2-")
    assert [Path(f).name for f in res["files"]] == ["m_p1.pdf", "m_p2-3.pdf"]
    assert run("pdf-documents", "fields", tmp_path / "form.pdf")[1]["fields"][0]["name"] == "name"
    (tmp_path / "v.json").write_text('{"name": "Luphia"}')
    assert run("pdf-documents", "fill", tmp_path / "form.pdf", tmp_path / "v.json", tmp_path / "f.pdf")[0] == 0
    assert run("pdf-documents", "fields", tmp_path / "f.pdf")[1]["fields"][0]["value"] == "Luphia"
    code, res = run("pdf-documents", "text", tmp_path / "two.pdf", "--pages", "2")
    assert "page two" in res["pages"][0]["text"]
    (tmp_path / "bad.json").write_text('{"nosuch": 1}')
    code, res = run("pdf-documents", "fill", tmp_path / "form.pdf", tmp_path / "bad.json", tmp_path / "g.pdf")
    assert code == 2 and "no such fields" in res["error"]


# ---------- data visualization ----------

def test_palette_is_colour_vision_safe_and_scripts_are_linted(tmp_path):
    code, res = run("data-visualization", "validate", "#0072B2,#D55E00,#009E73,#CC79A7,#222222,#666666")
    assert code == 0, res["problems"]
    code, res = run("data-visualization", "validate", "#E69F00,#F0E442")
    assert code == 1 and any("contrast" in p for p in res["problems"])
    s = tmp_path / "plot.py"
    s.write_text("import matplotlib.pyplot as plt\nplt.plot([1,2])\nplt.title('熱功率')\nplt.show()\n",
                 encoding="utf-8")
    code, res = run("data-visualization", "check", s)
    text = "\n".join(res["problems"])
    assert code == 1 and "CJK font" in text and "without savefig" in text


def test_sbplot_refuses_boxes_for_missing_glyphs(tmp_path, monkeypatch):
    need("matplotlib")
    sys.path.insert(0, str(SK / "data-visualization" / "scripts"))
    import sbplot
    monkeypatch.setattr(sbplot, "cjk_font", lambda: None)
    fig, ax = sbplot.figure()
    ax.set_title("熱功率")
    with pytest.raises(RuntimeError, match="missing from the font"):
        sbplot.save(fig, str(tmp_path / "x.png"))


# ---------- frontend ----------

def test_frontend_check_finds_phone_overflow_contrast_and_network(tmp_path):
    need("playwright")
    bad = tmp_path / "bad.html"
    bad.write_text('<html><head><script src="https://cdn.example.org/x.js"></script></head><body>'
                   '<p style="color:#bbb">淺灰色</p><div style="width:900px">寬</div></body></html>', encoding="utf-8")
    try:
        code, res = run("frontend-design", "check", bad)
    except json.JSONDecodeError:
        pytest.skip("no browser")
    if res.get("status") == "error":
        pytest.skip(res["error"])
    text = "\n".join(res["problems"])
    assert code == 1 and "sideways on a phone" in text and "contrast" in text and "network" in text
    assert "lang attribute" in text and "viewport" in text


# ---------- code review ----------

def test_code_review_scope_and_secrets(tmp_path):
    if not shutil.which("git"):
        pytest.skip("git missing")
    repo = tmp_path / "r"
    repo.mkdir()
    g = ["git", "-C", str(repo), "-c", "user.name=t", "-c", "user.email=t@t"]
    subprocess.run(g[:3] + ["init", "-q", "-b", "main"], check=True)
    (repo / "app.py").write_text("def f(x):\n    return x\n")
    subprocess.run(g + ["add", "."], check=True)
    subprocess.run(g + ["commit", "-qm", "init"], check=True)
    subprocess.run(g + ["checkout", "-qb", "feat"], check=True)
    (repo / "app.py").write_text('def f(x):\n    return x / 0\nAPI_KEY = "sk-ant-abcdefghijklmnopqrstuvwx"\n'
                                 'example = "password = \'changeme123\'"\n')
    subprocess.run(g + ["commit", "-qam", "feat"], check=True)
    code, res = run("code-review", "scope", repo)
    assert code == 0 and res["files"][0]["file"] == "app.py" and res["commits"][0].endswith("feat")
    assert "return x / 0" in res["diff"]
    code, res = run("code-review", "secrets", repo)
    assert code == 1 and len(res["problems"]) == 1 and res["problems"][0].startswith("app.py:3")
    assert "abcdefghijklmnop" not in res["problems"][0]           # the secret itself is not echoed


# ---------- visual check sees document pages ----------

@pytest.mark.skipif(not (SOFFICE and shutil.which("pdftoppm")), reason="LibreOffice / poppler not installed")
def test_visual_check_renders_document_pages(tmp_path):
    from sinobee.agent import visual
    docx = _docx(tmp_path)
    shots = visual._render_doc(docx)
    assert shots and shots[0]["label"] == "r.docx 第 1 頁" and shots[0]["image"][:4] == b"\x89PNG"
    assert zipfile.is_zipfile(docx)


def test_word_check_outside_sinobee_reads_table_text(tmp_path, monkeypatch):
    """Without SinoBee importable (the skill copied elsewhere) the check used to read paragraphs only, so a short
    document whose content is mostly a table was reported as empty."""
    import builtins
    import importlib.util
    docx = _docx(tmp_path)
    spec_ = importlib.util.spec_from_file_location("wd_run", SK / "word-documents" / "scripts" / "run.py")
    mod = importlib.util.module_from_spec(spec_)
    spec_.loader.exec_module(mod)
    real = builtins.__import__

    def no_sinobee(name, *a, **kw):
        if name.startswith("sinobee"):
            raise ImportError(name)
        return real(name, *a, **kw)
    monkeypatch.setattr(builtins, "__import__", no_sinobee)
    text = mod._text(docx)
    assert "功率 | 2 MWt" in text and len("".join(text.split())) >= 20
